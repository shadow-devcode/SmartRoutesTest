"""
Recalcula el desplazamiento de un Excel ya generado con distancias de carretera.

Los ruteros creados antes de que el motor consultara Mapbox llevan la
estimación en línea recta corregida por un factor: kilómetros parecidos a los
reales, pero tiempos que se quedaban a la mitad. Reprocesar el Excel entero
volvería a repartir los puntos y perdería los ajustes hechos a mano en el
calendario, así que este script toca SOLO las tres columnas del
desplazamiento: kilómetros, minutos entre sucursales y horarios de la jornada.

Consulta la misma API de Directions que dibuja la «Vista Carretera» del mapa,
con las paradas de cada día en su orden de visita.

Uso:
    python scripts/recalcular_carretera.py <excel> [--aplicar]

Sin `--aplicar` solo informa. Con él, guarda una copia `.bak` antes de tocar
nada y escribe de forma atómica.
"""
from __future__ import annotations

import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from dotenv import load_dotenv

load_dotenv(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".env"))

import pandas as pd

from route_engine.config import MAPBOX_ACCESS_TOKEN
from route_engine.excel_writer import recalcular_tramos_por_carretera
from utils.excel_atomic import escritura_atomica

HOJA = "Horarios_Detalle"


def recalcular(path: str, aplicar: bool) -> int:
    if not MAPBOX_ACCESS_TOKEN or not MAPBOX_ACCESS_TOKEN.startswith("pk."):
        print("[!] No hay MAPBOX_ACCESS_TOKEN en BACK/.env: sin él no se puede "
              "consultar la carretera.")
        return 0

    horarios = pd.read_excel(path, sheet_name=HOJA)
    # La fila de totales que arrastra la hoja no es una visita.
    mascara = horarios["Mercadista"].astype(str).str.strip().str.upper() != "TOTAL"
    visitas = horarios[mascara].copy()

    km_antes = visitas["kilometros entre sucurlas (km)"].sum()
    min_antes = visitas["Tiempo entre sucursal (min)"].sum()

    print(f"{path}")
    print(f"  antes : {km_antes:,.0f} km | {min_antes:,.0f} min de desplazamiento")
    print("  consultando la carretera de cada jornada...")
    actualizadas, total, tarde = recalcular_tramos_por_carretera(visitas)

    km_ahora = visitas["kilometros entre sucurlas (km)"].sum()
    min_ahora = visitas["Tiempo entre sucursal (min)"].sum()
    print(f"  ahora : {km_ahora:,.0f} km | {min_ahora:,.0f} min")
    print(f"  jornadas recalculadas: {actualizadas} de {total} "
          f"(las demás tienen una sola parada)")
    if tarde:
        print(f"  [!] {tarde} jornada(s) terminan pasada la hora de cierre con los "
              f"tiempos reales; no se recortan, conviene revisarlas.")

    if not actualizadas or not aplicar:
        if actualizadas:
            print("  (simulación: vuelve a lanzarlo con --aplicar para escribirlo)")
        return actualizadas

    shutil.copy2(path, path + ".bak")

    # Se escriben SOLO las tres celdas de cada visita, con openpyxl. Volcar el
    # libro entero con pandas dejaría en texto plano los anchos, los autofiltros
    # y las fórmulas SUBTOTAL de las otras hojas.
    from openpyxl import load_workbook

    libro = load_workbook(path)
    hoja = libro[HOJA]
    cabeceras = {
        str(celda.value).strip(): celda.column
        for celda in hoja[1]
        if celda.value is not None
    }
    columnas = ("kilometros entre sucurlas (km)", "Tiempo entre sucursal (min)", "Horario")
    faltan = [c for c in columnas if c not in cabeceras]
    if faltan:
        print(f"  [!] La hoja no tiene las columnas {faltan}; no se escribe nada.")
        return 0

    for indice, fila in visitas.iterrows():
        # +2: la fila 1 es la cabecera y openpyxl cuenta desde 1.
        fila_excel = int(indice) + 2
        for columna in columnas:
            hoja.cell(row=fila_excel, column=cabeceras[columna]).value = fila[columna]

    with escritura_atomica(path) as destino:
        libro.save(destino)
    print(f"  escrito. Copia previa en {path}.bak")
    return actualizadas


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        raise SystemExit(1)
    recalcular(sys.argv[1], "--aplicar" in sys.argv[2:])
